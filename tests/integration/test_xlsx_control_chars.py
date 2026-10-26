"""Sprint 97 / C: Steuerzeichen im Excel-Export als Symbol statt Abbruch.

openpyxl lehnt `\\x00`–`\\x08`, `\\x0b`, `\\x0c`, `\\x0e`–`\\x1f` in Zellen ab
(`IllegalCharacterError`). Seit Sprint 96 bleiben `\\v`, `\\f`, `\\x1c`–`\\x1e` im
Zellwert – ein Export solcher Zeilen brach ab. Ersetzt wird ausschließlich beim
Export (`io/_xlsx_safe.py`); DB-Werte bleiben unverändert.
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from sampling_tool import __version__
from sampling_tool.core.formatting import format_optional_timestamp
from sampling_tool.core.models import (
    Dataset,
    DatasetRow,
    Engagement,
    SampleConfig,
    SampleResult,
    SamplingMethod,
)
from sampling_tool.io._xlsx_safe import CONTROL_CHAR_PICTURES, safe_cell_value
from sampling_tool.io.exporter import ExcelExporter
from sampling_tool.persistence.database import Database
from sampling_tool.persistence.repositories import DatasetRepo, EngagementRepo

pytestmark = pytest.mark.integration

_EXPORT_NOW = datetime(2026, 10, 26, 9, 30, 0)
_DRAWN_AT = datetime(2026, 10, 25, 8, 0, 0, tzinfo=UTC)
_TEXT_COLUMN = "Buchungstext"


@pytest.fixture
def db() -> Iterator[Database]:
    database = Database(Path(":memory:"))
    database.migrate()
    yield database
    database.close()


def _export(db: Database, texts: list[str], out_dir: Path) -> tuple[Path, Dataset, DatasetRepo]:
    """Exportiert eine Stichprobe über alle Zeilen; ``texts`` ist die Textspalte."""
    eng = EngagementRepo(db.connect()).get_or_create(
        Engagement(
            auditor_name="Anna",
            client_name="ACME",
            auditor_position="Senior",
            audit_type="ISAE 3402",
        )
    )
    assert eng.id is not None
    repo = DatasetRepo(db.connect())
    rows = tuple(
        DatasetRow(row_id=i, values={"Nr": i, _TEXT_COLUMN: text})
        for i, text in enumerate(texts, start=1)
    )
    dataset = repo.create(
        Dataset(
            name="Buchungen",
            columns=("Nr", _TEXT_COLUMN),
            source_file="buchungen.csv",
            engagement_id=eng.id,
        ),
        rows,
    )
    sample = SampleResult(
        config=SampleConfig(method=SamplingMethod.SIMPLE, size=len(rows), seed=7),
        selected_row_ids=tuple(r.row_id for r in rows),
        population_size=len(rows),
        drawn_at=_DRAWN_AT,
        created_by="tester",
        id=1,
    )
    with patch("sampling_tool.io.exporter.datetime") as fake_datetime:
        fake_datetime.now.return_value = _EXPORT_NOW
        path = ExcelExporter().export_sample(
            sample,
            dataset,
            repo,
            columns=["Nr", _TEXT_COLUMN],
            output_dir=out_dir,
            custom_name="steuerzeichen",
            custom_id="1",
            engagement=eng,
            now=_EXPORT_NOW,
        )
    return path, dataset, repo


def _sheet_cells(path: Path, sheet: str) -> list[tuple[Any, ...]]:
    wb = load_workbook(path)
    try:
        return [tuple(c.value for c in row) for row in wb[sheet].iter_rows()]
    finally:
        wb.close()


def _sample_texts(path: Path) -> list[Any]:
    return [row[1] for row in _sample_rows(path)]


def _sample_rows(path: Path) -> list[tuple[Any, ...]]:
    return _sheet_cells(path, "Sample")[1:]


def _metadata_snapshot(
    path: Path,
) -> tuple[list[tuple[Any, ...]], list[str | None], dict[str, Any]]:
    """Werte, Datentypen und Spaltenbreiten des Metadaten-Blatts."""
    wb = load_workbook(path)
    try:
        ws = wb["Metadaten"]
        values = [tuple(c.value for c in row) for row in ws.iter_rows()]
        types = [c.data_type for row in ws.iter_rows() for c in row]
        widths = {key: dim.width for key, dim in ws.column_dimensions.items()}
        return values, types, widths
    finally:
        wb.close()


# Eingefroren auf `main` (e4c6fdd) mit derselben Eingabe ohne Steuerzeichen.
# Zeitstempel/App-Version werden wie im Exporter berechnet, damit die Referenz
# unabhängig von Zeitzone und Versionssprung des Laufs gilt.
def _frozen_metadata() -> list[tuple[Any, ...]]:
    return [
        ("Feld", "Wert"),
        ("Erstellt am", _EXPORT_NOW),
        ("Dataset", "Buchungen"),
        ("Quelldatei", "buchungen.csv"),
        ("Quellblatt", "—"),
        ("Kopfzeile", "—"),
        ("Dataset-ID", "1"),
        ("Sampling-Methode", "Einfach"),
        ("Angeforderte Größe", "2"),
        ("Tatsächliche Größe", "2"),
        ("Population (Zeilen)", "2"),
        ("Datensatz gesamt", "2"),
        ("Seed", "7"),
        ("Filter", "—"),
        ("Cluster-Feld", "—"),
        ("Schicht-Feld", "—"),
        ("Übergeordnete Stichprobe", "—"),
        ("Ableitung", "—"),
        ("Algorithmus-Version", "bdo-v1"),
        ("App-Version", __version__),
        ("Erstellt von", "tester"),
        ("Gezogen am", format_optional_timestamp(_DRAWN_AT)),
        ("Auditor", "Anna"),
        ("Auditor-Position", "Senior"),
        ("Mandant", "ACME"),
        ("Prüfungstyp", "ISAE 3402"),
    ]


class TestControlCharactersInExport:
    def test_each_illegal_char_mapped_to_symbol(self, db: Database, tmp_path: Path) -> None:
        illegal = [*range(0x00, 0x09), 0x0B, 0x0C, *range(0x0E, 0x20)]
        assert sorted(CONTROL_CHAR_PICTURES) == illegal
        texts = [f"a{chr(code)}b" for code in illegal]

        path, _dataset, _repo = _export(db, texts, tmp_path)

        assert _sample_texts(path) == [f"a{chr(0x2400 + code)}b" for code in illegal]

    def test_tab_newline_cr_unchanged(self, db: Database, tmp_path: Path) -> None:
        texts = ["Spalte\tWert", "Zeile 1\nZeile 2", "Ende\r"]
        assert [safe_cell_value(t) for t in texts] == texts

        path, _dataset, _repo = _export(db, texts, tmp_path)

        # Der XML-Parser beim Wiederöffnen macht aus CR ein LF (XML-Zeilenend-
        # Normalisierung, auf `main` genauso) – geprüft wird: kein Symbol.
        exported = _sample_texts(path)
        assert exported[:2] == texts[:2]
        assert not any(chr(0x2400 + c) in t for t in exported for c in (0x09, 0x0A, 0x0D))

    def test_formula_neutralization_still_applies(self, db: Database, tmp_path: Path) -> None:
        path, _dataset, _repo = _export(db, ["=\x0bSUMME(1)"], tmp_path)

        wb = load_workbook(path)
        try:
            cell = wb["Sample"].cell(row=2, column=2)
            assert cell.value == "=␋SUMME(1)"
            assert cell.data_type != "f"
        finally:
            wb.close()

    def test_metadata_note_only_when_replaced(self, db: Database, tmp_path: Path) -> None:
        clean_dir = tmp_path / "ohne"
        clean_dir.mkdir()
        clean, _dataset, _repo = _export(db, ["Miete Mai", "Strom"], clean_dir)
        values, types, widths = _metadata_snapshot(clean)
        assert values == _frozen_metadata()
        assert types == ["s"] * 2 + ["s", "d"] + ["s"] * (2 * (len(values) - 2))
        assert widths == {"A": 26.0, "B": 21.0}

        with_chars_dir = tmp_path / "mit"
        with_chars_dir.mkdir()
        other = Database(Path(":memory:"))
        try:
            other.migrate()
            noisy, _dataset, _repo = _export(
                other, ["Miete\x0bMai", "\x1c\x1dStrom"], with_chars_dir
            )
        finally:
            other.close()
        noisy_values = _metadata_snapshot(noisy)[0]
        assert noisy_values[:-1] == _frozen_metadata()
        assert noisy_values[-1] == (
            "Hinweis",
            "3 Steuerzeichen wurden als Symbol dargestellt (z. B. ␋).",
        )

    def test_db_values_unchanged(self, db: Database, tmp_path: Path) -> None:
        path, dataset, repo = _export(db, ["Miete\x0bMai"], tmp_path)

        assert _sample_texts(path) == ["Miete␋Mai"]
        assert dataset.id is not None
        stored = repo.get_rows_by_ids(dataset.id, [1])
        assert stored[0].values[_TEXT_COLUMN] == "Miete\x0bMai"
