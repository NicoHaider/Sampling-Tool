"""Filteranzeige über die echten Pfade und gespeicherte Filterwerte (Sprint 93).

Jede Fläche (Provenienz, AuditTrail-Details, Stichproben-Export, Projekt-Excel,
HTML, PDF) zeigt dieselbe Zeichenfolge. Gespeichert wird wie vorher – nur
Datums-/Zeitwerte in `filters` tragen jetzt ihren Typ (Typ-Tag wie bei Vorlagen
und Datensatz-Zellen), damit sie nach dem Neuöffnen nicht als Text erscheinen.
"""

from __future__ import annotations

import html
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from sampling_tool import __version__
from sampling_tool.audit.logger import AuditLogger
from sampling_tool.core.formatting import format_audit_details
from sampling_tool.core.models import (
    AuditEvent,
    Dataset,
    DatasetRow,
    Engagement,
    FilterCondition,
    FilterOperator,
    SampleConfig,
    SampleResult,
    SamplingMethod,
)
from sampling_tool.core.provenance import SamplingProvenance, format_event_details
from sampling_tool.core.sampling import create_sampler
from sampling_tool.io.exporter import ExcelExporter
from sampling_tool.io.html_report import HtmlReportGenerator
from sampling_tool.io.multi_report_exporter import MultiSheetReportExporter
from sampling_tool.io.pdf_report import AuditTrailPDF
from sampling_tool.persistence.database import Database
from sampling_tool.persistence.repositories import (
    AuditRepo,
    DatasetRepo,
    EngagementRepo,
    SampleRepo,
)
from sampling_tool.persistence.sample_repo import _filters_from_json

pytestmark = pytest.mark.integration

EQ, GTE, LT = FilterOperator.EQ, FilterOperator.GTE, FilterOperator.LT

VERTRIEB = FilterCondition("Kostenstelle", EQ, "Vertrieb")
AB_1000 = FilterCondition("Betrag", GTE, 1000)
BIS_20 = FilterCondition("Datum", LT, datetime(2024, 1, 20))
EXPECTED_TEXT = "Kostenstelle = „Vertrieb“ UND Betrag ≥ 1000 UND Datum < 2024-01-20"


def _rows() -> tuple[DatasetRow, ...]:
    return tuple(
        DatasetRow(
            row_id=i,
            values={
                "Kostenstelle": "Vertrieb" if i % 2 else "1.000",
                "Betrag": i * 100,
                "Datum": datetime(2024, 1, 1 + i % 28),
            },
        )
        for i in range(1, 61)
    )


class _Project:
    """Projekt-DB mit einer über den Produktpfad gezogenen und geloggten Stichprobe."""

    def __init__(self, tmp_path: Path, filters: tuple[FilterCondition, ...]) -> None:
        self.db = Database(tmp_path / "p.db")
        self.db.migrate()
        self.engagement = EngagementRepo(self.db.connect()).get_or_create(
            Engagement(auditor_name="Anna", client_name="ACME", audit_type="ISAE 3402")
        )
        assert self.engagement.id is not None
        self.dataset = DatasetRepo(self.db.connect()).create(
            Dataset(
                name="d",
                columns=("Kostenstelle", "Betrag", "Datum"),
                engagement_id=self.engagement.id,
            ),
            _rows(),
        )
        assert self.dataset.id is not None
        config = SampleConfig(method=SamplingMethod.SIMPLE, size=5, seed=4711, filters=filters)
        self.drawn = create_sampler(config).sample(
            DatasetRepo(self.db.connect()).iter_rows(self.dataset.id),
            population_size=self.dataset.row_count,
        )
        self.sample_id = SampleRepo(self.db.connect()).create_from_result(
            self.drawn, self.dataset.id
        )
        AuditLogger(AuditRepo(self.db.connect()), "Anna", self.engagement.id).log_sampling(
            self.drawn, self.sample_id, self.dataset.id, dataset_rows=self.dataset.row_count
        )

    def stored(self) -> tuple[str, str | None, str]:
        conn = self.db.connect()
        (details,) = conn.execute(
            "SELECT details_json FROM audit_events WHERE event_type = 'sampling'"
        ).fetchone()
        filter_value, filters = conn.execute(
            "SELECT filter_value, filters FROM samples WHERE id = ?", (self.sample_id,)
        ).fetchone()
        return details, filter_value, filters

    def reloaded_sample(self) -> SampleResult:
        sample = SampleRepo(self.db.connect()).get_by_id(self.sample_id)
        assert sample is not None
        return sample

    def events(self) -> list[AuditEvent]:
        assert self.engagement.id is not None
        return AuditRepo(self.db.connect()).list_for_engagement(self.engagement.id)

    def sampling_event(self) -> AuditEvent:
        return next(e for e in self.events() if e.event_type == "sampling")


@pytest.fixture
def project(tmp_path: Path) -> Iterator[_Project]:
    proj = _Project(tmp_path, (VERTRIEB, AB_1000, BIS_20))
    yield proj
    proj.db.close()


def _provenance(sample: SampleResult) -> SamplingProvenance:
    return SamplingProvenance.from_sample_result(sample, dataset_id=1, app_version=__version__)


def _pdf_text(path: Path) -> str:
    return " ".join(" ".join(page.extract_text().split()) for page in PdfReader(path).pages)


class TestFilterDisplayAllPaths:
    def test_provenance_filter_text(self, project: _Project) -> None:
        assert _provenance(project.drawn).filter_text == EXPECTED_TEXT
        assert _provenance(project.reloaded_sample()).filter_text == EXPECTED_TEXT
        details = format_event_details(project.sampling_event())
        assert f"Filter: {EXPECTED_TEXT}" in details

    def test_legacy_event_single_filter_keys(self) -> None:
        details = {
            "method": "simple",
            "filter_field": "Kostenstelle",
            "filter_value": "Vertrieb",
            "filter_operator": "eq",
        }
        assert (
            format_audit_details(details) == "Methode: Einfach · Filter: Kostenstelle = „Vertrieb“"
        )
        legacy_date = {
            "filter_field": "Datum",
            "filter_value": "2024-06-30T00:00:00",
            "filter_operator": "lte",
        }
        assert format_audit_details(legacy_date) == "Filter: Datum ≤ 2024-06-30"
        legacy_none = {"filter_field": "Notiz", "filter_value": None, "filter_operator": "eq"}
        assert format_audit_details(legacy_none) == "Filter: Notiz = (kein Wert)"

    def test_audit_details_line(self) -> None:
        details = {
            "method": "simple",
            "filters": [
                {"field": "Konto", "operator": "eq", "value": "1.000"},
                {"field": "Konto", "operator": "ne", "value": 1000},
            ],
        }
        assert format_audit_details(details) == (
            "Methode: Einfach · Filter: Konto = „1.000“ UND Konto ≠ 1000"
        )

    def test_reports_show_same_text(self, project: _Project, tmp_path: Path) -> None:
        sample = project.reloaded_sample()
        expected = _provenance(sample).filter_text
        assert expected == EXPECTED_TEXT
        events = project.events()
        assert project.dataset.id is not None
        mapping = {project.sample_id: project.dataset.id}

        sample_xlsx = ExcelExporter().export_sample(
            sample,
            project.dataset,
            DatasetRepo(project.db.connect()),
            columns=["Kostenstelle"],
            output_dir=tmp_path,
            custom_name="x",
            custom_id="1",
        )
        meta = {
            r[0]: r[1]
            for r in load_workbook(sample_xlsx)["Metadaten"].iter_rows(2, values_only=True)
        }
        assert meta["Filter"] == expected

        report = MultiSheetReportExporter().export(
            project.engagement,
            [project.dataset],
            [sample],
            events,
            tmp_path / "r.xlsx",
            dataset_ids_by_sample=mapping,
        )
        cells = [
            str(c)
            for ws in load_workbook(report).worksheets
            for row in ws.iter_rows(values_only=True)
            for c in row
            if c is not None
        ]
        assert any(c == expected for c in cells), "Samples-Blatt"
        assert any(f"Filter: {expected}" in c for c in cells), "AuditTrail-Blatt"

        page = HtmlReportGenerator().render(
            project.engagement,
            [project.dataset],
            [sample],
            events,
            tmp_path / "r.html",
            include_charts=False,
            dataset_ids_by_sample=mapping,
        )
        text = html.unescape(page.read_text(encoding="utf-8"))
        assert expected in text
        assert f"Filter: {expected}" in text

        pdf = AuditTrailPDF(briefpapier=None).render(project.engagement, events, tmp_path / "a.pdf")
        assert f"Filter: {expected}" in _pdf_text(pdf)


# Referenz vor Sprint 93 (Stand c3dbf4a) – eingefroren, nicht neu berechnen.
_REFERENCE_DETAILS_JSON = (
    '{"dataset_rows":3,"population_basis":"auswahl","dataset_id":1,"method":"simple",'
    '"size_requested":1,"filters":[{"field":"K","operator":"eq","value":"Vertrieb"},'
    '{"field":"K","operator":"eq","value":"1.000"},{"field":"N","operator":"gte","value":1000},'
    '{"field":"F","operator":"lt","value":5000.5},{"field":"B","operator":"eq","value":true}],'
    '"cluster_field":null,"stratum_field":null,"stratify_mode":"proportional",'
    '"parent_sample_id":null,"parent_relation":null,"algorithm_version":"bdo-v1",'
    '"app_version":"<APP>","created_by":"system"}'
)
_REFERENCE_SAMPLES_FILTERS = (
    '[{"field":"K","operator":"eq","value":"Vertrieb"},{"field":"K","operator":"eq","value":"1.000"},'
    '{"field":"N","operator":"gte","value":1000},{"field":"F","operator":"lt","value":5000.5},'
    '{"field":"B","operator":"eq","value":true}]'
)
_REFERENCE_EMPTY_NONE = (
    '[{"field":"K","operator":"eq","value":""},{"field":"K","operator":"eq","value":null}]'
)
_REFERENCE_SINGLE_DETAILS = (
    '"filters":[{"field":"K","operator":"eq","value":"1.000"}],'
    '"filter_field":"K","filter_value":"1.000","filter_operator":"eq",'
)
# Stand c3dbf4a: gleicher Seed, gleiche Filter (siehe `project`).
_REFERENCE_SELECTION = (29, 37, 39, 43, 59)


def _store(tmp_path: Path, filters: tuple[FilterCondition, ...]) -> tuple[str, str | None, str]:
    """Wie der Produktpfad, aber über einen kleinen Datensatz (Referenz-Shape)."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    db = Database(tmp_path / "s.db")
    db.migrate()
    try:
        eng = EngagementRepo(db.connect()).get_or_create(
            Engagement(auditor_name="A", client_name="C")
        )
        assert eng.id is not None
        ds = DatasetRepo(db.connect()).create(
            Dataset(name="d", columns=("K",), engagement_id=eng.id),
            tuple(DatasetRow(row_id=i, values={"K": "V"}) for i in range(1, 4)),
        )
        assert ds.id is not None
        result = SampleResult(
            config=SampleConfig(method=SamplingMethod.SIMPLE, size=1, seed=1, filters=filters),
            selected_row_ids=(1,),
            population_size=3,
        )
        sample_id = SampleRepo(db.connect()).create_from_result(result, ds.id)
        AuditLogger(AuditRepo(db.connect()), "A", eng.id).log_sampling(
            result, sample_id, ds.id, dataset_rows=3
        )
        conn = db.connect()
        (details,) = conn.execute("SELECT details_json FROM audit_events").fetchone()
        filter_value, stored = conn.execute("SELECT filter_value, filters FROM samples").fetchone()
        return details, filter_value, stored
    finally:
        db.close()


class TestNoStorageChange:
    def test_stored_event_details_unchanged(self, tmp_path: Path) -> None:
        filters = (
            FilterCondition("K", EQ, "Vertrieb"),
            FilterCondition("K", EQ, "1.000"),
            FilterCondition("N", GTE, 1000),
            FilterCondition("F", LT, 5000.5),
            FilterCondition("B", EQ, True),
        )
        details, filter_value, stored = _store(tmp_path / "a", filters)
        assert details == _REFERENCE_DETAILS_JSON.replace("<APP>", __version__)
        assert filter_value is None
        assert stored == _REFERENCE_SAMPLES_FILTERS

        empty_none = (FilterCondition("K", EQ, ""), FilterCondition("K", EQ, None))
        _, _, stored = _store(tmp_path / "b", empty_none)
        assert stored == _REFERENCE_EMPTY_NONE

        details, filter_value, stored = _store(tmp_path / "c", (FilterCondition("K", EQ, "1.000"),))
        assert _REFERENCE_SINGLE_DETAILS in details
        assert filter_value == '"1.000"'

    def test_sample_selection_unchanged(self, project: _Project) -> None:
        assert project.drawn.selected_row_ids == _REFERENCE_SELECTION
        assert project.reloaded_sample().selected_row_ids == _REFERENCE_SELECTION


class TestTypedTemporalFilters:
    """Option 3: Datums-/Zeitwerte in `filters` tragen ihren Typ (Sprint 93)."""

    _TAGGED = '{"__type__":"datetime","v":"2024-01-20T00:00:00"}'

    def test_filters_store_type_tag(self, tmp_path: Path) -> None:
        details, _, stored = _store(tmp_path, (BIS_20,))
        assert f'"filters":[{{"field":"Datum","operator":"lt","value":{self._TAGGED}}}]' in details
        assert stored == f'[{{"field":"Datum","operator":"lt","value":{self._TAGGED}}}]'

    def test_legacy_single_keys_unchanged(self, tmp_path: Path) -> None:
        details, filter_value, _ = _store(tmp_path, (BIS_20,))
        assert '"filter_value":"2024-01-20T00:00:00"' in details
        assert filter_value == '"2024-01-20T00:00:00"'

    def test_reload_keeps_datetime(self, project: _Project) -> None:
        reloaded = project.reloaded_sample().config.filters
        assert reloaded == project.drawn.config.filters
        assert type(reloaded[2].value) is datetime

    def test_text_same_before_and_after_reopen(self, project: _Project) -> None:
        assert _provenance(project.reloaded_sample()).filter_text == (
            _provenance(project.drawn).filter_text
        )

    def test_aware_datetime_roundtrip(self, tmp_path: Path) -> None:
        aware = FilterCondition("Datum", GTE, datetime(2024, 1, 20, 8, 30, tzinfo=UTC))
        _, _, stored = _store(tmp_path, (aware,))
        assert '"v":"2024-01-20T08:30:00+00:00"' in stored
        assert _filters_from_json(stored) == (aware,)
