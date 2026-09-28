"""Sprint 89 / B: Stichprobennummer im PDF, Population = Auswahlgrundlage in allen Berichten.

Alte Filter-Stichproben (Sampling-Event ohne `population_basis`) behalten ihre
gespeicherte Population und bekommen den Altbestand-Hinweis; neue zeigen
zusätzlich „Datensatz gesamt".
"""

from __future__ import annotations

import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from sampling_tool.audit.logger import AuditLogger
from sampling_tool.config import LEGACY_FILTER_POPULATION_NOTE
from sampling_tool.core.models import (
    AuditEvent,
    Dataset,
    DatasetRow,
    Engagement,
    FilterOperator,
    SampleConfig,
    SampleResult,
    SamplingMethod,
)
from sampling_tool.core.provenance import (
    format_event_details,
    population_predates_basis,
    population_text,
)
from sampling_tool.io.exporter import ExcelExporter
from sampling_tool.io.html_report import HtmlReportGenerator
from sampling_tool.io.multi_report_exporter import MultiSheetReportExporter
from sampling_tool.io.pdf_report import _EVENT_TABLE_HEADER, _build_event_table
from sampling_tool.persistence.database import Database
from sampling_tool.persistence.repositories import (
    AuditRepo,
    DatasetRepo,
    EngagementRepo,
    SampleRepo,
)

pytestmark = pytest.mark.integration

_WHEN = datetime(2026, 9, 27, 10, 0, tzinfo=UTC)
_NEW_DETAILS: dict[str, Any] = {
    "dataset_rows": 500,
    "population_basis": "auswahl",
    "method": "simple",
    "filter_field": "Kostenstelle",
    "filter_value": "Vertrieb",
    "filter_operator": "eq",
}
_OLD_DETAILS: dict[str, Any] = {
    "method": "simple",
    "filter_field": "Kostenstelle",
    "filter_value": "Vertrieb",
    "filter_operator": "eq",
}


def _sampling_event(sample_id: int, total: int, details: dict[str, Any]) -> AuditEvent:
    return AuditEvent(
        event_type="sampling",
        engagement_id=1,
        user_name="anna",
        timestamp=_WHEN,
        sample_id=sample_id,
        sample_size=5,
        sample_percent=5 / total * 100,
        total_count=total,
        seed=7,
        details=details,
        id=sample_id,
    )


def _filtered_sample(sample_id: int, population: int) -> SampleResult:
    return SampleResult(
        config=SampleConfig(
            method=SamplingMethod.SIMPLE,
            size=5,
            seed=7,
            filter_field="Kostenstelle",
            filter_operator=FilterOperator.EQ,
            filter_value="Vertrieb",
        ),
        selected_row_ids=(1, 2, 3, 4, 5),
        population_size=population,
        drawn_at=_WHEN,
        id=sample_id,
    )


@pytest.fixture
def engagement() -> Engagement:
    return Engagement(auditor_name="Anna", client_name="ACME", audit_type="ISAE 3402", id=1)


class TestPopulationText:
    def test_new_population_is_plain(self) -> None:
        assert population_text(45, legacy_filter=False) == "45"

    def test_legacy_filter_population_gets_the_note(self) -> None:
        assert population_text(500, legacy_filter=True) == (
            f"500 ({LEGACY_FILTER_POPULATION_NOTE})"
        )
        assert "Stand vor Sprint 89" in LEGACY_FILTER_POPULATION_NOTE

    def test_missing_basis_or_event_predates(self) -> None:
        assert population_predates_basis(None) is True
        assert population_predates_basis(_OLD_DETAILS) is True
        assert population_predates_basis(_NEW_DETAILS) is False


class TestEventDetails:
    def test_new_sampling_event_leads_with_population_and_dataset_total(self) -> None:
        text = format_event_details(_sampling_event(10, 45, _NEW_DETAILS))
        assert text.startswith("Population: 45 · Datensatz gesamt: 500")
        assert "Population bezogen auf" not in text
        assert LEGACY_FILTER_POPULATION_NOTE not in text

    def test_old_filter_event_carries_the_note(self) -> None:
        text = format_event_details(_sampling_event(1, 500, _OLD_DETAILS))
        assert text.startswith(f"Population: 500 ({LEGACY_FILTER_POPULATION_NOTE})")

    def test_old_unfiltered_event_has_no_note(self) -> None:
        text = format_event_details(_sampling_event(1, 500, {"method": "simple"}))
        assert text.startswith("Population: 500 · ")
        assert LEGACY_FILTER_POPULATION_NOTE not in text

    def test_other_events_are_unchanged(self) -> None:
        evt = AuditEvent(event_type="reset", engagement_id=1, details={"dataset_id": 3})
        assert format_event_details(evt) == "Datensatz-ID: 3"


class TestPdfSampleColumn:
    def test_header_and_cell(self) -> None:
        evt = _sampling_event(10, 45, _NEW_DETAILS)
        (table,) = _build_event_table([evt])
        column = _EVENT_TABLE_HEADER.index("Stichprobe")
        assert table._cellvalues[0][column] == "Stichprobe"
        assert table._cellvalues[1][column] == "#10"

    def test_event_without_sample_shows_dash(self) -> None:
        evt = AuditEvent(event_type="import", engagement_id=1, timestamp=_WHEN, id=1)
        (table,) = _build_event_table([evt])
        assert table._cellvalues[1][_EVENT_TABLE_HEADER.index("Stichprobe")] == "—"

    def test_legacy_note_in_the_pdf(self, tmp_path: Path, engagement: Engagement) -> None:
        from pypdf import PdfReader

        from sampling_tool.io.pdf_report import AuditTrailPDF

        out = AuditTrailPDF(briefpapier=None).render(
            engagement, [_sampling_event(1, 500, _OLD_DETAILS)], tmp_path / "a.pdf"
        )
        text = " ".join(" ".join(p.extract_text().split()) for p in PdfReader(str(out)).pages)
        assert "gesamter Datensatz" in text
        assert "Stand vor Sprint 89" in text


class TestExcelReport:
    def _export(self, tmp_path: Path, engagement: Engagement) -> Any:
        dataset = Dataset(name="d", columns=("Kostenstelle",), engagement_id=1, row_count=500, id=3)
        out = MultiSheetReportExporter().export(
            engagement,
            [dataset],
            [_filtered_sample(1, 500), _filtered_sample(10, 45)],
            [_sampling_event(1, 500, _OLD_DETAILS), _sampling_event(10, 45, _NEW_DETAILS)],
            tmp_path / "r.xlsx",
            dataset_ids_by_sample={1: 3, 10: 3},
            sampling_details_by_sample={1: _OLD_DETAILS, 10: _NEW_DETAILS},
        )
        return load_workbook(out)

    def test_samples_sheet_population_and_dataset_total(
        self, tmp_path: Path, engagement: Engagement
    ) -> None:
        ws = self._export(tmp_path, engagement)["3. Samples"]
        header = [c.value for c in ws[1]]
        rows = {r[0]: dict(zip(header, r, strict=True)) for r in ws.iter_rows(2, values_only=True)}
        assert rows[10]["Population"] == 45
        assert rows[10]["Datensatz gesamt"] == 500
        assert rows[1]["Population"] == f"500 ({LEGACY_FILTER_POPULATION_NOTE})"
        assert rows[1]["Datensatz gesamt"] == 500

    def test_audit_trail_details_start_with_population(
        self, tmp_path: Path, engagement: Engagement
    ) -> None:
        ws = self._export(tmp_path, engagement)["2. AuditTrail"]
        header = [c.value for c in ws[1]]
        details = [r[header.index("Details")] for r in ws.iter_rows(2, values_only=True)]
        assert details[0].startswith(f"Population: 500 ({LEGACY_FILTER_POPULATION_NOTE})")
        assert details[1].startswith("Population: 45 · Datensatz gesamt: 500")

    def test_without_mapping_no_note(self, tmp_path: Path, engagement: Engagement) -> None:
        out = MultiSheetReportExporter().export(
            engagement, [], [_filtered_sample(1, 500)], [], tmp_path / "r.xlsx"
        )
        ws = load_workbook(out)["3. Samples"]
        header = [c.value for c in ws[1]]
        row = next(ws.iter_rows(2, values_only=True))
        assert row[header.index("Population")] == 500
        assert row[header.index("Datensatz gesamt")] == "—"


class TestHtmlReport:
    def test_samples_table(self, tmp_path: Path, engagement: Engagement) -> None:
        dataset = Dataset(name="d", columns=("Kostenstelle",), engagement_id=1, row_count=500, id=3)
        out = HtmlReportGenerator().render(
            engagement,
            [dataset],
            [_filtered_sample(1, 500), _filtered_sample(10, 45)],
            [_sampling_event(1, 500, _OLD_DETAILS), _sampling_event(10, 45, _NEW_DETAILS)],
            tmp_path / "r.html",
            include_charts=False,
            dataset_ids_by_sample={1: 3, 10: 3},
            sampling_details_by_sample={1: _OLD_DETAILS, 10: _NEW_DETAILS},
        )
        html = out.read_text(encoding="utf-8")
        section = html.split("<h2>Stichproben</h2>", 1)[1].split("</table>", 1)[0]
        headers = re.findall(r"<th>(.*?)</th>", section)
        rows = {}
        for row in re.findall(r"<tr>\s*(<td>#.*?)</tr>", section, re.S):
            cells = re.findall(r"<td>(.*?)</td>", row, re.S)
            rows[cells[0]] = dict(zip(headers, cells, strict=True))
        assert rows["#10"]["Population"] == "45"
        assert rows["#10"]["Datensatz gesamt"] == "500"
        assert rows["#1"]["Population"] == f"500 ({LEGACY_FILTER_POPULATION_NOTE})"
        assert "Population: 45 · Datensatz gesamt: 500" in html


class TestSampleExportMetadata:
    def _meta(self, tmp_path: Path, *, predates: bool) -> dict[str, Any]:
        db = Database(tmp_path / "p.db")
        db.migrate()
        try:
            eng = EngagementRepo(db.connect()).get_or_create(
                Engagement(auditor_name="Anna", client_name="ACME", audit_type="ISAE 3402")
            )
            ds = DatasetRepo(db.connect()).create(
                Dataset(name="d", columns=("Kostenstelle",), engagement_id=eng.id),
                tuple(DatasetRow(row_id=i, values={"Kostenstelle": "V"}) for i in range(1, 8)),
            )
            out = ExcelExporter().export_sample(
                _filtered_sample(1, 7),
                ds,
                DatasetRepo(db.connect()),
                columns=["Kostenstelle"],
                output_dir=tmp_path,
                custom_name="x",
                custom_id="1",
                population_predates_basis=predates,
            )
        finally:
            db.close()
        ws = load_workbook(out)["Metadaten"]
        return {r[0]: r[1] for r in ws.iter_rows(2, values_only=True)}

    def test_dataset_total_row(self, tmp_path: Path) -> None:
        meta = self._meta(tmp_path, predates=False)
        assert meta["Population (Zeilen)"] == "7"
        assert meta["Datensatz gesamt"] == "7"

    def test_legacy_note(self, tmp_path: Path) -> None:
        meta = self._meta(tmp_path, predates=True)
        assert meta["Population (Zeilen)"] == f"7 ({LEGACY_FILTER_POPULATION_NOTE})"


class TestSamplingDetailsBySample:
    def test_reads_every_sampling_event(self, db: Database, engagement_id: int) -> None:
        ds = DatasetRepo(db.connect()).create(
            Dataset(name="d", columns=("a",), engagement_id=engagement_id),
            (DatasetRow(row_id=1, values={"a": 1}),),
        )
        assert ds.id is not None
        result = SampleResult(
            config=SampleConfig(method=SamplingMethod.SIMPLE, size=1, seed=1),
            selected_row_ids=(1,),
            population_size=1,
        )
        sample_id = SampleRepo(db.connect()).create_from_result(result, ds.id)
        AuditLogger(AuditRepo(db.connect()), "anna", engagement_id).log_sampling(
            result, sample_id, ds.id, dataset_rows=1
        )
        mapping = AuditRepo(db.connect()).sampling_details_by_sample(engagement_id)
        assert mapping[sample_id]["population_basis"] == "auswahl"
        assert mapping[sample_id]["dataset_rows"] == 1


class TestExcelPercentIsANumber:
    """Sprint 89 / D2: „%" ist in allen Blättern eine Zahl mit Prozentformat."""

    def test_both_sheets(self, tmp_path: Path, engagement: Engagement) -> None:
        out = MultiSheetReportExporter().export(
            engagement,
            [],
            [_filtered_sample(10, 45)],
            [_sampling_event(10, 45, _NEW_DETAILS)],
            tmp_path / "r.xlsx",
        )
        wb = load_workbook(out)
        for sheet, column in (("2. AuditTrail", "%"), ("3. Samples", "Anteil %")):
            ws = wb[sheet]
            header = [c.value for c in ws[1]]
            cell = ws.cell(row=2, column=header.index(column) + 1)
            assert isinstance(cell.value, float), sheet
            assert cell.value == pytest.approx(5 / 45)
            assert "%" in cell.number_format, sheet
