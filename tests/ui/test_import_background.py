"""Sprint 97: Vorprüfung und Vorschau des Imports laufen im Hintergrund.

Kern (`TestNoParseOnUiThread`): `CalamineWorkbook.get_sheet_by_name` ist die
Funktion, die ein Blatt komplett einliest (Inventur Sprint 97: `from_path` und
`sheet_names` lesen keine Zellen, `height`/`width` sind nach dem Laden gratis).
Der Test zeichnet bei jedem Aufruf den Thread auf – auf dem UI-Thread darf kein
Blatt mehr geladen werden.
"""

from __future__ import annotations

import sqlite3
import threading
import time
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from PyQt6.QtWidgets import QDialog, QMessageBox
from pytestqt.qtbot import QtBot
from python_calamine import CalamineError, CalamineWorkbook

from sampling_tool.core.cancellation import CancellationToken
from sampling_tool.core.models import Engagement
from sampling_tool.io import import_preflight
from sampling_tool.io import importer as importer_module
from sampling_tool.io.importer import ExcelImporter, ImportProbe
from sampling_tool.persistence.database import Database
from sampling_tool.persistence.repositories import EngagementRepo
from sampling_tool.ui.controllers._factories import default_import_options_factory
from sampling_tool.ui.controllers.main_controller import MainController
from sampling_tool.ui.dialogs.progress_dialog import wait_for_detached_workers
from sampling_tool.ui.main_window import MainWindow
from sampling_tool.ui.recent import RecentEngagementsStore
from sampling_tool.ui.workers.task_worker import ProgressReporter, _ProgressEmitter
from sampling_tool.ui.workers.tasks import ExcelImportTask
from tests.integration.test_import_result_unchanged import (
    _make_rich_xlsx,
    _reference_values_to_json,
)
from tests.ui._progress_cancel import CancelClicker

pytestmark = pytest.mark.ui

_WC = "sampling_tool.ui.controllers.workspace_controller"


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _stub_post_import_id_dialog(monkeypatch: pytest.MonkeyPatch) -> None:
    """Der optionale ID-Spalten-Dialog nach dem Import ist modal (Sprint 31)."""
    monkeypatch.setattr(
        "sampling_tool.ui.dialogs.id_column_dialog.IdColumnDialog.exec",
        lambda _self: int(QDialog.DialogCode.Rejected),
    )


@pytest.fixture
def window(qtbot: QtBot) -> MainWindow:
    win = MainWindow()
    qtbot.addWidget(win)
    return win


@pytest.fixture
def engagement_db(tmp_path: Path) -> Path:
    db_path = tmp_path / "engagement.db"
    db = Database(db_path)
    db.migrate()
    EngagementRepo(db.connect()).get_or_create(
        Engagement(
            auditor_name="Anna",
            client_name="ACME",
            auditor_position="Senior",
            audit_type="ISAE 3402",
        )
    )
    db.close()
    return db_path


@pytest.fixture
def clean_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "sauber.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Daten"
    ws.append(["Konto", "Betrag"])
    for i in range(1, 6):
        ws.append([f"K{i}", i * 100])
    wb.save(path)
    return path


@pytest.fixture
def three_sheet_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "drei_blaetter.xlsx"
    wb = Workbook()
    first = wb.active
    assert first is not None
    first.title = "Erstes"
    first.append(["a", "b"])
    first.append([1, 2])
    second = wb.create_sheet("Zweites")
    second.append(["x", "y", "z"])
    second.append([10, 20, 30])
    third = wb.create_sheet("Drittes")
    third.append(["Bericht"])
    third.append([None])
    third.append(["Konto", "Betrag"])
    third.append(["K1", 5])
    wb.save(path)
    return path


@pytest.fixture
def clean_csv(tmp_path: Path) -> Path:
    path = tmp_path / "sauber.csv"
    path.write_bytes(b"Konto;Betrag\nK1;100\nK2;200\n")
    return path


@pytest.fixture
def controller(window: MainWindow, tmp_path: Path) -> Iterator[MainController]:
    ctrl = MainController(window, recent_store=RecentEngagementsStore(path=tmp_path / "r.json"))
    yield ctrl
    ctrl.engagement.handle_close_engagement()


def _controller_with_dialog(
    window: MainWindow, tmp_path: Path, factory: Callable[..., Any]
) -> MainController:
    return MainController(
        window,
        recent_store=RecentEngagementsStore(path=tmp_path / "r.json"),
        import_options_dialog_factory=factory,
    )


class _SheetLoadRecorder:
    """Zeichnet jeden `get_sheet_by_name`-Aufruf mit Thread auf."""

    def __init__(self) -> None:
        self.calls: list[tuple[str, bool]] = []

    def install(self, monkeypatch: pytest.MonkeyPatch) -> None:
        original = CalamineWorkbook.get_sheet_by_name
        recorder = self

        def _recording(wb: CalamineWorkbook, name: str) -> Any:
            recorder.calls.append((name, threading.current_thread() is threading.main_thread()))
            return original(wb, name)

        monkeypatch.setattr(CalamineWorkbook, "get_sheet_by_name", _recording)

    @property
    def on_main_thread(self) -> list[str]:
        return [name for name, main in self.calls if main]


def _run_import(controller: MainController, path: Path) -> None:
    with (
        patch(f"{_WC}.QFileDialog.getOpenFileName", return_value=(str(path), "")),
        patch(f"{_WC}.QMessageBox.information"),
    ):
        controller.workspace.handle_import_excel()


def _dataset_count(db_path: Path) -> int:
    conn = sqlite3.connect(db_path)
    try:
        return int(conn.execute("SELECT COUNT(*) FROM datasets").fetchone()[0])
    finally:
        conn.close()


def _import_event_count(db_path: Path) -> int:
    conn = sqlite3.connect(db_path)
    try:
        row = conn.execute("SELECT COUNT(*) FROM audit_events WHERE event_type = 'import'")
        return int(row.fetchone()[0])
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# Kern: kein Parse auf dem UI-Thread
# ---------------------------------------------------------------------------


class TestNoParseOnUiThread:
    def test_clean_single_sheet_import_parses_off_main_thread(
        self,
        controller: MainController,
        engagement_db: Path,
        clean_xlsx: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        controller.engagement.handle_open_engagement(engagement_db)
        recorder = _SheetLoadRecorder()
        recorder.install(monkeypatch)

        _run_import(controller, clean_xlsx)

        assert _dataset_count(engagement_db) == 1
        assert recorder.on_main_thread == []
        assert len(recorder.calls) == 2, recorder.calls

    def test_multi_sheet_dialog_switching_does_not_parse(
        self,
        window: MainWindow,
        tmp_path: Path,
        engagement_db: Path,
        three_sheet_xlsx: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        switches: list[str] = []

        def factory(*args: Any) -> Any:
            dialog = default_import_options_factory(*args)

            def _exec() -> int:
                combo = dialog._sheet_combo
                assert combo is not None
                for index in (1, 2, 0):
                    combo.setCurrentIndex(index)
                    switches.append(str(combo.currentData()))
                combo.setCurrentIndex(1)
                dialog._on_accept()
                return int(QDialog.DialogCode.Accepted)

            dialog.exec = _exec  # type: ignore[method-assign]
            return dialog

        controller = _controller_with_dialog(window, tmp_path, factory)
        try:
            controller.engagement.handle_open_engagement(engagement_db)
            recorder = _SheetLoadRecorder()
            recorder.install(monkeypatch)

            _run_import(controller, three_sheet_xlsx)
        finally:
            controller.engagement.handle_close_engagement()

        assert switches == ["Zweites", "Drittes", "Erstes"]
        assert _dataset_count(engagement_db) == 1
        assert recorder.on_main_thread == []
        # Ein Laden je Blatt in der Probe, eins im Import (gewähltes Blatt).
        assert sorted(name for name, _ in recorder.calls) == sorted(
            ["Erstes", "Zweites", "Drittes", "Zweites"]
        )

    def test_csv_probe_off_main_thread(
        self,
        controller: MainController,
        engagement_db: Path,
        clean_csv: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        controller.engagement.handle_open_engagement(engagement_db)
        opens: list[bool] = []
        real_open = open

        def _recording_open(file: Any, *args: Any, **kwargs: Any) -> Any:
            if Path(file) == clean_csv:
                opens.append(threading.current_thread() is threading.main_thread())
            return real_open(file, *args, **kwargs)

        monkeypatch.setattr(importer_module, "open", _recording_open, raising=False)

        _run_import(controller, clean_csv)

        assert _dataset_count(engagement_db) == 1
        assert opens, "Der Importer muss die CSV öffnen"
        assert not any(opens), "Keine CSV-Öffnung auf dem UI-Thread"


# ---------------------------------------------------------------------------
# Probe == bisherige Einzelfunktionen
# ---------------------------------------------------------------------------

_FIXTURE_CSV_DIR = Path(__file__).resolve().parent.parent / "fixtures" / "csv"

# xlsx-Fixtures der bestehenden Import-Tests (conftest) plus die dieses Moduls.
_XLSX_FIXTURES = (
    "simple_xlsx",
    "multi_sheet_xlsx",
    "empty_xlsx",
    "leading_blank_xlsx",
    "duplicate_columns_xlsx",
    "clean_xlsx",
    "three_sheet_xlsx",
    "ambiguous_xlsx",
)


@pytest.fixture
def ambiguous_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "zahlen.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append([100, 200, 300])
    ws.append([400, 500, 600])
    wb.save(path)
    return path


def _probe_sources() -> list[str]:
    return [*_XLSX_FIXTURES, *(f"csv:{p.name}" for p in sorted(_FIXTURE_CSV_DIR.glob("*.csv")))]


class TestProbeEquivalence:
    @pytest.mark.parametrize("source", _probe_sources())
    def test_probe_matches_legacy_functions(
        self, request: pytest.FixtureRequest, source: str
    ) -> None:
        importer = ExcelImporter()
        if source.startswith("csv:"):
            path = _FIXTURE_CSV_DIR / source.removeprefix("csv:")
            probe = importer.probe(path)
            assert probe.sheets == ()
            assert dict(probe.previews) == {path.name: importer.preview_csv(path)}
        else:
            path = request.getfixturevalue(source)
            probe = importer.probe(path)
            sheets = importer.list_sheets(path)
            assert list(probe.sheets) == sheets
            assert dict(probe.previews) == {
                info.name: importer.preview_sheet(path, info.name) for info in sheets
            }
        assert probe.needs_dialog == importer.requires_options_dialog(path)
        assert probe.dimension_reject is None
        assert probe.dimension_warnings == ()

    def test_dimension_messages_unchanged(
        self, tmp_path: Path, three_sheet_xlsx: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Wortlaut wie `_check_dimensions` vor Sprint 97 (dort eingefroren)."""
        big = tmp_path / "gross.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Daten"
        ws.append(["Konto", "Betrag"])
        for i in range(1_000):
            ws.append([f"K{i}", i])
        wb.save(big)

        monkeypatch.setattr(import_preflight, "WARN_IMPORT_ROWS", 1_000)
        probe = ExcelImporter().probe(big)
        assert probe.dimension_reject is None
        assert probe.dimension_warnings == ("Sheet 'Daten' hat 1,001 Zeilen.",)

        monkeypatch.setattr(import_preflight, "WARN_IMPORT_ROWS", 1)
        probe = ExcelImporter().probe(three_sheet_xlsx)
        assert probe.dimension_reject is None
        assert probe.dimension_warnings == (
            "Sheet 'Erstes' hat 2 Zeilen.",
            "Sheet 'Zweites' hat 2 Zeilen.",
            "Sheet 'Drittes' hat 4 Zeilen.",
        )

        monkeypatch.setattr(import_preflight, "MAX_IMPORT_ROWS", 999)
        probe = ExcelImporter().probe(big)
        assert probe.dimension_reject == "Sheet 'Daten' hat zu viele Zeilen (1,001, erlaubt: 999)."

        monkeypatch.setattr(import_preflight, "MAX_IMPORT_COLUMNS", 2)
        probe = ExcelImporter().probe(three_sheet_xlsx)
        assert probe.dimension_reject == "Sheet 'Zweites' hat zu viele Spalten (3, erlaubt: 2)."
        # Wie vorher: Abbruch am ersten auffälligen Blatt, Warnungen bis dahin.
        assert [info.name for info in probe.sheets] == ["Erstes", "Zweites"]
        assert probe.dimension_warnings == ("Sheet 'Erstes' hat 2 Zeilen.",)

    def test_unreadable_workbook_message_unchanged(
        self, clean_xlsx: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def _boom(_path: str) -> None:
            raise CalamineError("kaputt")

        monkeypatch.setattr(CalamineWorkbook, "from_path", staticmethod(_boom))

        probe = ExcelImporter().probe(clean_xlsx)

        assert probe == ImportProbe(
            dimension_reject="Excel-Datei konnte nicht gelesen werden: sauber.xlsx (kaputt)."
        )


# ---------------------------------------------------------------------------
# Ablauf im Controller
# ---------------------------------------------------------------------------


class TestProbeFlow:
    def test_cancel_during_probe_aborts_import(
        self,
        controller: MainController,
        window: MainWindow,
        engagement_db: Path,
        clean_xlsx: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        controller.engagement.handle_open_engagement(engagement_db)
        original_probe = ExcelImporter.probe
        waited: list[bool] = []

        def _waiting_probe(self: ExcelImporter, path: Path, max_rows: int = 20) -> ImportProbe:
            # Läuft im Worker: wartet, bis „Abbrechen" das Token setzt.
            assert self.cancellation is not None
            deadline = time.monotonic() + 10
            while not self.cancellation.is_set():
                assert time.monotonic() < deadline, "Abbruch kam nie an"
            waited.append(True)
            return original_probe(self, path, max_rows)

        monkeypatch.setattr(ExcelImporter, "probe", _waiting_probe)
        # Der echte Knopf: `QProgressDialog.cancel()` allein sendet `canceled`
        # nicht und setzte das Token nie. Sprint 98: der Dialog erscheint erst
        # nach 300 ms, der Helfer wartet darauf.
        clicker = CancelClicker()
        try:
            with patch(f"{_WC}.QMessageBox.warning") as mock_warning:
                _run_import(controller, clean_xlsx)
        finally:
            clicker.stop()

        assert clicker.labels == ["Datei wird geprüft…"]
        # Sprint 98 / C: die Probe liest nur – „Abbrechen" wartet nicht auf den
        # Worker, der läuft abgelöst zu Ende. Erst danach steht `waited` fest.
        wait_for_detached_workers()
        assert waited == [True], "Das Token muss im Worker angekommen sein"
        mock_warning.assert_not_called()
        assert _dataset_count(engagement_db) == 0
        assert _import_event_count(engagement_db) == 0

    def test_warnings_shown_once_combined(
        self,
        controller: MainController,
        engagement_db: Path,
        clean_xlsx: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        controller.engagement.handle_open_engagement(engagement_db)
        monkeypatch.setattr(import_preflight, "WARN_IMPORT_FILE_SIZE_BYTES", 10)
        monkeypatch.setattr(import_preflight, "WARN_IMPORT_ROWS", 1)

        with patch(
            f"{_WC}.QMessageBox.question", return_value=QMessageBox.StandardButton.No
        ) as mock_question:
            _run_import(controller, clean_xlsx)

        mock_question.assert_called_once()
        assert mock_question.call_args.args[1:3] == (
            "Große Datei",
            "Die gewählte Datei ist ungewöhnlich groß oder umfangreich:\n\n"
            "Datei ist groß (0 MB).\n"
            "Sheet 'Daten' hat 6 Zeilen."
            "\n\nImport trotzdem fortsetzen?",
        )
        assert _dataset_count(engagement_db) == 0
        assert _import_event_count(engagement_db) == 0

    def test_reject_from_probe_shows_error(
        self,
        controller: MainController,
        engagement_db: Path,
        clean_xlsx: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        controller.engagement.handle_open_engagement(engagement_db)
        monkeypatch.setattr(import_preflight, "MAX_IMPORT_COLUMNS", 1)

        with (
            patch(f"{_WC}.QMessageBox.warning") as mock_warning,
            patch(f"{_WC}.QMessageBox.question") as mock_question,
        ):
            _run_import(controller, clean_xlsx)

        mock_warning.assert_called_once()
        assert mock_warning.call_args.args[2] == (
            "Import abgelehnt: Sheet 'Daten' hat zu viele Spalten (2, erlaubt: 1)."
        )
        mock_question.assert_not_called()
        assert _dataset_count(engagement_db) == 0

    def test_reject_from_preflight_skips_probe(
        self,
        controller: MainController,
        engagement_db: Path,
        tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Ohne ZIP-Signatur erreicht die Datei weder Probe noch calamine."""
        controller.engagement.handle_open_engagement(engagement_db)
        fake = tmp_path / "kein_excel.xlsx"
        fake.write_text("kein Excel", encoding="utf-8")
        probes: list[Path] = []
        monkeypatch.setattr(
            ExcelImporter, "probe", lambda _self, path, max_rows=20: probes.append(path)
        )

        with patch(f"{_WC}.QMessageBox.warning") as mock_warning:
            _run_import(controller, fake)

        assert probes == []
        mock_warning.assert_called_once()
        assert mock_warning.call_args.args[2] == (
            "Import abgelehnt: Datei 'kein_excel.xlsx' hat keine gültige Excel-Signatur "
            "(ZIP-Container erwartet)."
        )

    def test_import_result_unchanged_through_new_flow(
        self, controller: MainController, engagement_db: Path, tmp_path: Path
    ) -> None:
        path = _make_rich_xlsx(tmp_path)
        controller.engagement.handle_open_engagement(engagement_db)

        _run_import(controller, path)

        flow = _stored_import(engagement_db)
        # Referenz 1: eingefrorener Encoder aus test_import_result_unchanged.py.
        expected_rows = list(ExcelImporter().import_file(path).rows)
        assert [r[0] for r in flow["rows"]] == [r.row_id for r in expected_rows]
        assert [r[1] for r in flow["rows"]] == [
            _reference_values_to_json(r.values) for r in expected_rows
        ]
        # Referenz 2: derselbe Import direkt über die Task, ohne den neuen Ablauf.
        direct_db = tmp_path / "direkt.db"
        _create_engagement_db(direct_db)
        emitter = _ProgressEmitter()
        ExcelImportTask(
            path=path,
            db_path=direct_db,
            engagement_id=1,
            user_name=str(flow["event"][3]),
        ).run(ProgressReporter(emitter), CancellationToken())
        assert flow == _stored_import(direct_db)


def _create_engagement_db(db_path: Path) -> None:
    db = Database(db_path)
    db.migrate()
    EngagementRepo(db.connect()).get_or_create(
        Engagement(
            auditor_name="Anna",
            client_name="ACME",
            auditor_position="Senior",
            audit_type="ISAE 3402",
        )
    )
    db.close()


def _stored_import(db_path: Path) -> dict[str, Any]:
    """Alles, was ein Import persistiert – ohne Zeitstempel."""
    conn = sqlite3.connect(db_path)
    try:
        dataset = conn.execute(
            "SELECT id, name, source_file, row_count, columns_json FROM datasets"
        ).fetchall()
        rows = conn.execute(
            "SELECT row_index, values_json FROM dataset_rows ORDER BY row_index"
        ).fetchall()
        event = conn.execute(
            "SELECT event_type, import_file, total_count, user_name, details_json "
            "FROM audit_events WHERE event_type = 'import'"
        ).fetchone()
        return {"dataset": dataset, "rows": rows, "event": event}
    finally:
        conn.close()
