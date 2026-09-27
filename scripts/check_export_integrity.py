"""Bestandsprüfung (Sprint 86): enthalten Sample-Exporte Zeilen des richtigen Datensatzes?

Bis Sprint 86 konnte „Sample exportieren" nach dem Wiederöffnen eines Projekts
die Zeilen-IDs einer Stichprobe auf einen ANDEREN Datensatz anwenden. Die Datei
enthielt dann dessen Zeilen, und ihr Blatt „Metadaten" nannte dessen
Dataset-ID. Dieses Skript vergleicht für jedes `export`-Event mit Stichprobe die
Dataset-ID aus der exportierten Datei mit dem Datensatz, auf dem die
Stichprobe tatsächlich gezogen wurde (`samples.dataset_id`).

Rein lesend: Projekt- und Exportdateien werden nie geändert oder repariert.
Ohne WAL wird die Projektdatei mit `mode=ro&immutable=1` geöffnet (legt keine
`-wal`/`-shm`-Dateien an). Liegt daneben eine nicht leere `-wal`-Datei – die
App ist offen oder hat den WAL noch nicht zurückgeschrieben –, stünden neuere
Export-Events nur dort; `immutable` würde sie ignorieren und „0 geprüft"
melden (Sprint 87 / D). Dann werden `.db`, `-wal` und ggf. `-shm` in ein
temporäres Verzeichnis kopiert und die KOPIE normal geöffnet.

Aufruf (auch bei offener App):
    python scripts/check_export_integrity.py "<Pfad>/<Mandant>.db"

Ausgabe je Event: OK / ABWEICHUNG / DATEI FEHLT / UNKLAR.
Exit-Code:
    0  alles OK oder nichts zu prüfen (auch: keine Audit-Events; UNKLAR und
       DATEI FEHLT werden nur in der Zusammenfassung gezählt)
    1  mindestens eine ABWEICHUNG
    2  Projektdatei nicht lesbar (nicht gefunden, keine SQLite-Datei, oder
       z. B. nur als Cloud-Platzhalter vorhanden)
"""

from __future__ import annotations

import argparse
import json
import shutil
import sqlite3
import sys
import tempfile
from collections.abc import Iterator
from contextlib import closing, contextmanager
from dataclasses import dataclass
from pathlib import Path

from _script_io import force_utf8_stdout
from openpyxl import load_workbook

try:
    from enum import StrEnum
except ImportError:  # pragma: no cover – nur auf Python < 3.11 erreichbar
    sys.exit(
        "check_export_integrity.py braucht Python 3.11 oder neuer "
        f"(gefunden: {sys.version.split()[0]})."
    )

_META_SHEET = "Metadaten"
_DATASET_ID_LABEL = "Dataset-ID"


class ExportStatus(StrEnum):
    """Ergebnis je `export`-Event."""

    OK = "OK"
    MISMATCH = "ABWEICHUNG"
    FILE_MISSING = "DATEI FEHLT"
    UNCLEAR = "UNKLAR"


@dataclass(frozen=True, slots=True)
class ExportCheck:
    """Prüfergebnis für ein `export`-Event mit Stichprobe."""

    event_id: int
    timestamp: str
    sample_id: int
    file: Path
    sample_dataset_id: int | None
    file_dataset_id: int | None
    status: ExportStatus
    note: str = ""


@dataclass(frozen=True, slots=True)
class _ExportEvent:
    event_id: int
    timestamp: str
    sample_id: int | None
    file: Path
    archived_previous: Path | None


class ProjectUnreadableError(Exception):
    """Die Projektdatei ließ sich nicht lesen (I/O, Cloud-Platzhalter)."""


@dataclass(frozen=True, slots=True)
class ProjectCheck:
    """Ergebnis für eine Projektdatei."""

    checks: list[ExportCheck]
    wal_applied: bool
    has_audit_trail: bool


def check_exports(db_path: Path) -> list[ExportCheck]:
    """Prüft alle Sample-Exporte der Projektdatei, in Event-Reihenfolge."""
    return check_project(db_path).checks


def check_project(db_path: Path) -> ProjectCheck:
    """Wie `check_exports`, plus: wurde ein WAL berücksichtigt, gibt es Audit-Events?

    Raises:
        ProjectUnreadableError: Datei oder WAL nicht lesbar (`OSError`,
            `sqlite3.OperationalError`).
        sqlite3.DatabaseError: die Datei ist keine SQLite-Datenbank.
    """
    wal_applied = _has_pending_wal(db_path)
    try:
        with _open_project(db_path, wal_applied) as conn:
            if not _has_table(conn, "audit_events"):
                return ProjectCheck(checks=[], wal_applied=wal_applied, has_audit_trail=False)
            events = _export_events(conn)
            owners = {
                int(row["id"]): int(row["dataset_id"])
                for row in conn.execute("SELECT id, dataset_id FROM samples")
            }
    except (OSError, sqlite3.OperationalError) as exc:
        raise ProjectUnreadableError(str(exc)) from exc

    locations = _current_locations(events)
    checks: list[ExportCheck] = []
    for event in events:
        if event.sample_id is None:
            continue  # Berichts-Export (PDF/Excel-/HTML-Bericht), keine Stichprobe.
        checks.append(_check_one(event, event.sample_id, locations[event.event_id], owners))
    return ProjectCheck(checks=checks, wal_applied=wal_applied, has_audit_trail=True)


def _sidecar(db_path: Path, suffix: str) -> Path:
    return db_path.with_name(db_path.name + suffix)


def _has_pending_wal(db_path: Path) -> bool:
    wal = _sidecar(db_path, "-wal")
    return wal.is_file() and wal.stat().st_size > 0


@contextmanager
def _open_project(db_path: Path, with_wal: bool) -> Iterator[sqlite3.Connection]:
    if not with_wal:
        with closing(_connect_read_only(db_path)) as conn:
            yield conn
        return
    # Die Kopie wendet den WAL beim Öffnen an; das Original (auch das einer
    # gerade offenen App) bleibt byte-genau, es entsteht dort kein `-shm`.
    with tempfile.TemporaryDirectory(prefix="export-check-") as tmp:
        copy = Path(tmp) / db_path.name
        shutil.copy2(db_path, copy)
        for suffix in ("-wal", "-shm"):
            sidecar = _sidecar(db_path, suffix)
            if sidecar.is_file():
                shutil.copy2(sidecar, _sidecar(copy, suffix))
        # Vor dem Aufräumen schließen – Windows löscht keine offenen Dateien.
        with closing(sqlite3.connect(copy)) as conn:
            conn.row_factory = sqlite3.Row
            yield conn


def _connect_read_only(db_path: Path) -> sqlite3.Connection:
    # Wie `persistence/db_preflight.py`: `mode=ro` allein legt neben einer
    # WAL-Datei trotzdem `-wal`/`-shm` an; `immutable=1` nicht.
    uri = f"{db_path.resolve().as_uri()}?mode=ro&immutable=1"
    conn = sqlite3.connect(uri, uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def _has_table(conn: sqlite3.Connection, name: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?", (name,)
    ).fetchone()
    return row is not None


def _export_events(conn: sqlite3.Connection) -> list[_ExportEvent]:
    events: list[_ExportEvent] = []
    for row in conn.execute(
        "SELECT id, timestamp, sample_id, export_file, details_json FROM audit_events "
        "WHERE event_type = 'export' AND export_file IS NOT NULL ORDER BY id"
    ):
        details = json.loads(row["details_json"]) if row["details_json"] else {}
        archived = details.get("archived_previous") if isinstance(details, dict) else None
        events.append(
            _ExportEvent(
                event_id=int(row["id"]),
                timestamp=str(row["timestamp"]),
                sample_id=None if row["sample_id"] is None else int(row["sample_id"]),
                file=Path(row["export_file"]),
                archived_previous=Path(archived) if isinstance(archived, str) else None,
            )
        )
    return events


def _current_locations(events: list[_ExportEvent]) -> dict[int, Path]:
    """Wo liegt die Datei jedes Events heute?

    Überschreibt ein späterer Export dieselbe Datei, verschiebt die App die alte
    Fassung vorher (Sprint 84 / C) und vermerkt das Ziel am NEUEN Event. Die
    alte Fassung gehört aber zum VORHERIGEN Event an diesem Pfad.
    """
    locations: dict[int, Path] = {}
    holder_of: dict[Path, int] = {}
    for event in events:
        previous = holder_of.get(event.file)
        if previous is not None and event.archived_previous is not None:
            locations[previous] = event.archived_previous
        locations[event.event_id] = event.file
        holder_of[event.file] = event.event_id
    return locations


def _check_one(
    event: _ExportEvent, sample_id: int, file: Path, owners: dict[int, int]
) -> ExportCheck:
    sample_dataset_id = owners.get(sample_id)

    def result(status: ExportStatus, file_dataset_id: int | None, note: str = "") -> ExportCheck:
        return ExportCheck(
            event_id=event.event_id,
            timestamp=event.timestamp,
            sample_id=sample_id,
            file=file,
            sample_dataset_id=sample_dataset_id,
            file_dataset_id=file_dataset_id,
            status=status,
            note=note,
        )

    if not file.is_file():
        return result(ExportStatus.FILE_MISSING, None)
    try:
        file_dataset_id = _dataset_id_in_file(file)
    except Exception as exc:  # beschädigte/fremde Datei: melden, nicht abbrechen
        return result(ExportStatus.UNCLEAR, None, f"Datei nicht lesbar: {exc}")
    if file_dataset_id is None:
        return result(ExportStatus.UNCLEAR, None, "kein Feld „Dataset-ID“ im Blatt „Metadaten“")
    if sample_dataset_id is None:
        return result(ExportStatus.UNCLEAR, file_dataset_id, "Stichprobe nicht in der Projektdatei")
    if file_dataset_id != sample_dataset_id:
        return result(ExportStatus.MISMATCH, file_dataset_id)
    return result(ExportStatus.OK, file_dataset_id)


def _dataset_id_in_file(path: Path) -> int | None:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if _META_SHEET not in wb.sheetnames:
            return None
        for row in wb[_META_SHEET].iter_rows(values_only=True):
            if len(row) >= 2 and row[0] == _DATASET_ID_LABEL:
                value = row[1]
                try:
                    return int(str(value).strip())
                except ValueError:
                    return None  # „—": Export ohne dokumentierte Dataset-ID
        return None
    finally:
        wb.close()


def _format(check: ExportCheck) -> str:
    head = f"{check.status.value:<11} Event {check.event_id} ({check.timestamp}), "
    head += f"Stichprobe #{check.sample_id}"
    match check.status:
        case ExportStatus.OK:
            detail = f"Datensatz {check.sample_dataset_id}"
        case ExportStatus.MISMATCH:
            detail = (
                f"gezogen auf Datensatz {check.sample_dataset_id}, "
                f"Datei enthält Datensatz {check.file_dataset_id}"
            )
        case ExportStatus.FILE_MISSING:
            detail = "Datei nicht vorhanden"
        case ExportStatus.UNCLEAR:
            detail = check.note
    return f"{head}: {detail} – {check.file}"


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Prüft rein lesend, ob Sample-Exporte die Zeilen des Datensatzes enthalten, "
            "auf dem ihre Stichprobe gezogen wurde."
        )
    )
    parser.add_argument("db", type=Path, help="Projektdatei (.db)")
    return parser


def main(argv: list[str] | None = None) -> int:
    # Vor jeder Ausgabe: gedruckte Pfade können beliebige Unicode-Zeichen
    # enthalten, die cp1252 (Windows, umgeleitetes stdout) nicht kodieren kann.
    force_utf8_stdout()
    args = _build_parser().parse_args(argv)
    db_path: Path = args.db
    if not db_path.is_file():
        print(f"Projektdatei nicht gefunden: {db_path}", file=sys.stderr)
        return 2
    try:
        project = check_project(db_path)
    except ProjectUnreadableError as exc:
        print(
            f"Projektdatei nicht lesbar (evtl. nur in der Cloud, nicht heruntergeladen): "
            f"{db_path} ({exc})",
            file=sys.stderr,
        )
        return 2
    except sqlite3.DatabaseError as exc:
        print(f"Keine lesbare Projektdatei: {db_path} ({exc})", file=sys.stderr)
        return 2

    if project.wal_applied:
        print(
            "Hinweis: offene Änderungen aus der WAL-Datei wurden berücksichtigt (über eine Kopie)."
        )
        print()
    if not project.has_audit_trail:
        print("Nichts zu prüfen: keine Audit-Events in dieser Projektdatei.")
        return 0

    checks = project.checks
    for check in checks:
        print(_format(check))
    counts = {status: sum(c.status is status for c in checks) for status in ExportStatus}
    print()
    print(
        f"{len(checks)} Sample-Exporte geprüft: "
        + ", ".join(f"{counts[status]} {status.value}" for status in ExportStatus)
    )
    return 1 if counts[ExportStatus.MISMATCH] else 0


if __name__ == "__main__":
    sys.exit(main())
